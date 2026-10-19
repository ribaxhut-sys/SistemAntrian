"""Laporan rentang tanggal: angka, pembagian per loket/petugas, dan ekspor."""

from __future__ import annotations

from datetime import timedelta

import pytest

from app.core import clock
from app.core.errors import AuthorizationError, ValidationError
from tests.conftest import TEST_PASSWORD


@pytest.fixture
def laporan_hari_ini(auth, antrian_siap, cro_queue, cro_survey, terbitkan, layani_sampai_selesai):
    """Dua selesai (satu dinilai 5), satu tidak hadir, satu masih menunggu."""
    from app.services.report_service import ReportService

    pertama = layani_sampai_selesai("A")
    cro_survey.submit(pertama.id, 5)
    layani_sampai_selesai("A")
    terbitkan("A", 2)
    loket = antrian_siap["loket"][0]
    dipanggil = cro_queue.call_next(loket)
    cro_queue.no_show(dipanggil.id)

    auth.logout()
    auth.login("supervisor", TEST_PASSWORD)
    hari = clock.today()
    return ReportService(auth).laporan(hari, hari)


def test_angka_total(laporan_hari_ini):
    total = laporan_hari_ini.total
    assert total.terbit == 4
    assert total.selesai == 2
    assert total.tidak_hadir == 1
    assert total.dipanggil == 3
    assert total.jumlah_survey == 1
    assert total.rata_kepuasan == 5
    assert total.rata_layanan_detik is not None


def test_laporan_mencatat_prioritas_dan_yang_dicabut(auth, antrian_siap):
    from app.services.queue_service import QueueService
    from app.services.report_service import ReportService

    auth.logout()
    auth.mulai_sesi_kios()
    nomor = QueueService(auth).create_queue(
        antrian_siap["kategori"]["A"],
        perangkat="tablet-laporan",
        prioritas_mandiri=True,
    )
    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    QueueService(auth).cabut_prioritas(nomor.id)
    auth.logout()
    auth.login("supervisor", TEST_PASSWORD)

    laporan = ReportService(auth).laporan(clock.today(), clock.today())

    assert laporan.total.prioritas == 1
    assert laporan.total.prioritas_dicabut == 1
    assert laporan.per_hari[0].prioritas == 1
    assert laporan.riwayat[0].prioritas
    assert laporan.riwayat[0].prioritas_dicabut


def test_pembagian_per_loket_petugas_dan_hari(laporan_hari_ini):
    (loket,) = laporan_hari_ini.per_loket
    assert loket.selesai == 2
    assert loket.label_kepuasan == "5,0"
    (petugas,) = laporan_hari_ini.per_petugas
    assert petugas.dipanggil == 3
    (hari,) = laporan_hari_ini.per_hari
    assert hari.terbit == 4
    assert sum(b.terbit for b in laporan_hari_ini.per_layanan) == 4


def test_rentang_tanpa_data_kosong(auth, antrian_siap):
    from app.services.report_service import ReportService

    auth.logout()
    auth.login("supervisor", TEST_PASSWORD)
    kemarin = clock.today() - timedelta(days=1)
    laporan = ReportService(auth).laporan(kemarin, kemarin)
    assert laporan.total.terbit == 0
    assert laporan.per_hari == []


def test_rentang_tidak_sah_ditolak(auth, antrian_siap):
    from app.services.report_service import ReportService

    auth.logout()
    auth.login("supervisor", TEST_PASSWORD)
    hari = clock.today()
    layanan = ReportService(auth)
    with pytest.raises(ValidationError):
        layanan.laporan(hari, hari - timedelta(days=1))
    with pytest.raises(ValidationError):
        layanan.laporan(hari - timedelta(days=400), hari)


def test_hanya_supervisor(auth, antrian_siap):
    from app.services.report_service import ReportService

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    hari = clock.today()
    with pytest.raises(AuthorizationError):
        ReportService(auth).laporan(hari, hari)


def test_laporan_petugas_hanya_memuat_nomor_miliknya(
    auth, laporan_hari_ini
):
    from app.services.report_service import ReportService, angka_ringkasan

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    hari = clock.today()

    laporan = ReportService(auth).laporan_petugas(hari, hari)

    assert laporan.nama_petugas == "Petugas Cro"
    assert laporan.total.terbit == 3
    assert len(laporan.riwayat) == 3
    assert {baris.petugas for baris in laporan.riwayat} == {"Petugas Cro"}
    assert "Menunggu" not in {baris.status for baris in laporan.riwayat}
    kartu = [nama for nama, _ in angka_ringkasan(laporan)]
    assert "Rata-rata menunggu" not in kartu
    assert "Nomor dipanggil" in kartu


def test_admin_tidak_dapat_memakai_jalur_laporan_pribadi(auth, antrian_siap):
    from app.services.report_service import ReportService

    auth.logout()
    auth.login("admin", TEST_PASSWORD)
    hari = clock.today()

    with pytest.raises(AuthorizationError):
        ReportService(auth).laporan_petugas(hari, hari)


def test_ekspor_pdf_dan_excel(laporan_hari_ini, tmp_path):
    from openpyxl import load_workbook

    from app.services.report_export import ekspor_excel, ekspor_pdf

    pdf = ekspor_pdf(laporan_hari_ini, tmp_path / "lap.pdf", "KANTOR UJI")
    assert pdf.read_bytes().startswith(b"%PDF")

    xlsx = ekspor_excel(laporan_hari_ini, tmp_path / "lap.xlsx")
    buku = load_workbook(xlsx)
    assert buku.sheetnames == [
        "Ringkasan",
        "Per hari",
        "Per layanan",
        "Per loket",
        "Per petugas",
        "Petugas per layanan",
        "Riwayat pelayanan",
    ]
    loket = buku["Per loket"]
    assert loket["A1"].value == "Loket"
    assert loket["C2"].value == 2  # angka, bukan teks
    assert loket["G2"].value == 5

    rincian = buku["Petugas per layanan"]
    assert [sel.value for sel in rincian[1]][:4] == ["Petugas", "Jenis layanan", "Dipanggil", "Selesai"]
    assert rincian["D2"].value == 2
    assert rincian.auto_filter.ref

    riwayat = buku["Riwayat pelayanan"]
    assert riwayat.max_row == 1 + 4  # judul + satu baris per nomor
    assert riwayat["G2"].is_date  # jam terbit tersimpan sebagai waktu
    assert sorted(riwayat.cell(row=r, column=13).value or 0 for r in range(2, 6)) == [0, 0, 0, 5]


def test_ekspor_pribadi_tidak_memuat_tabel_petugas_lain(
    auth, laporan_hari_ini, tmp_path
):
    from openpyxl import load_workbook

    from app.services.report_export import ekspor_excel, ekspor_pdf, nama_berkas_bawaan
    from app.services.report_service import ReportService

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    hari = clock.today()
    laporan = ReportService(auth).laporan_petugas(hari, hari)

    assert nama_berkas_bawaan(laporan, "pdf").startswith("laporan-saya-petugas-cro-")
    pdf = ekspor_pdf(laporan, tmp_path / "saya.pdf", "KANTOR UJI")
    assert pdf.read_bytes().startswith(b"%PDF")
    xlsx = ekspor_excel(laporan, tmp_path / "saya.xlsx")
    buku = load_workbook(xlsx)
    assert buku["Ringkasan"]["A1"].value == "Laporan Pelayanan Saya"
    assert buku["Ringkasan"]["A3"].value == "Petugas Cro"
    assert "Per petugas" not in buku.sheetnames
    assert "Petugas per layanan" not in buku.sheetnames


def test_petugas_per_layanan_dan_riwayat(laporan_hari_ini):
    (pasangan,) = laporan_hari_ini.petugas_layanan
    total, rincian = pasangan
    assert total is laporan_hari_ini.per_petugas[0]
    (layanan,) = rincian
    assert layanan.kelompok == total.nama
    assert (layanan.dipanggil, layanan.selesai, layanan.tidak_hadir) == (3, 2, 1)
    assert layanan.jumlah_survey == 1

    riwayat = laporan_hari_ini.riwayat
    assert len(riwayat) == 4
    assert [t.terbit for t in riwayat] == sorted(t.terbit for t in riwayat)
    assert sorted(t.status for t in riwayat) == ["Menunggu", "Selesai", "Selesai", "Tidak Hadir"]
    selesai = [t for t in riwayat if t.status == "Selesai"]
    assert all(t.layanan_detik is not None for t in selesai)
    assert all(t.petugas == total.nama for t in selesai)
    assert sorted(t.nilai or 0 for t in selesai) == [0, 5]
    menunggu = [t for t in riwayat if t.dipanggil is None]
    assert len(menunggu) == 1 and menunggu[0].petugas == ""


def test_rentang_cepat():
    from datetime import date

    from app.ui.dashboard.report_page import rentang_cepat

    hari = date(2026, 10, 7)
    assert rentang_cepat("hari_ini", hari) == (hari, hari)
    assert rentang_cepat("7_hari", hari) == (date(2026, 10, 1), hari)
    assert rentang_cepat("bulan_ini", hari) == (date(2026, 10, 1), hari)
    assert rentang_cepat("bulan_lalu", hari) == (date(2026, 9, 1), date(2026, 9, 30))


def test_halaman_laporan_tampil(qtbot, laporan_hari_ini, auth):
    from app.services.report_service import ReportService
    from app.ui.dashboard.report_page import LaporanPage

    halaman = LaporanPage(ReportService(auth), "KANTOR")
    qtbot.addWidget(halaman)
    halaman._pakai_rentang("hari_ini")

    assert halaman.kartu[0].nilai.text() == "4"
    assert halaman.tabel["per_loket"].rowCount() == 1
    assert halaman.ekspor_pdf.isEnabled()


def test_halaman_laporan_pribadi_menyimpan_excel_secara_atomik(
    qtbot, laporan_hari_ini, auth, monkeypatch, tmp_path
):
    from app.services.report_service import ReportService
    from app.ui.dashboard.report_page import LaporanPage

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    halaman = LaporanPage(ReportService(auth), "KANTOR", pribadi=True)
    qtbot.addWidget(halaman)
    tujuan_tanpa_akhiran = tmp_path / "laporan-saya"
    monkeypatch.setattr(
        "app.ui.dashboard.report_page.QFileDialog.getSaveFileName",
        lambda *_: (str(tujuan_tanpa_akhiran), "Excel (*.xlsx)"),
    )

    halaman._ekspor("xlsx")

    tujuan = tujuan_tanpa_akhiran.with_suffix(".xlsx")
    qtbot.waitUntil(lambda: tujuan.is_file(), timeout=10_000)
    assert tujuan.is_file()
    assert not tujuan.with_name(f".{tujuan.name}.sementara").exists()
    assert tujuan.name in halaman.status_ekspor.text()


@pytest.mark.parametrize("peran", ["supervisor", "admin"])
def test_menu_laporan_untuk_supervisor_dan_admin(qtbot, auth, antrian_siap, peran):
    from app.ui.main_window import MainWindow

    auth.logout()
    auth.login(peran, TEST_PASSWORD)
    jendela = MainWindow(auth)
    qtbot.addWidget(jendela)
    menu = [jendela.navigasi.item(i).text() for i in range(jendela.navigasi.count())]
    assert "Laporan" in menu
    assert "Laporan Saya" not in menu


def test_menu_laporan_saya_hanya_untuk_cro(qtbot, auth, antrian_siap):
    from app.ui.main_window import MainWindow

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    jendela = MainWindow(auth)
    qtbot.addWidget(jendela)
    menu = [jendela.navigasi.item(i).text() for i in range(jendela.navigasi.count())]

    assert "Laporan Saya" in menu
    assert "Laporan" not in menu


def _audit_unduhan():
    import json

    import sqlalchemy as sa

    from app.core.database import session_scope
    from app.models import AuditAction, AuditLog

    with session_scope() as s:
        return [
            (user_id, json.loads(meta))
            for user_id, meta in s.execute(
                sa.select(AuditLog.user_id, AuditLog.meta)
                .where(AuditLog.action == AuditAction.EXPORT_REPORT.value)
                .order_by(AuditLog.id)
            )
        ]


def test_unduhan_laporan_tercatat_siapa_dan_cakupannya(auth, users, antrian_siap):
    from app.services.report_service import ReportService

    hari = clock.today()
    auth.logout()
    auth.login("admin", TEST_PASSWORD)
    layanan = ReportService(auth)
    layanan.catat_unduhan(layanan.laporan(hari, hari), "pdf")

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    layanan = ReportService(auth)
    layanan.catat_unduhan(layanan.laporan_petugas(hari, hari), "xlsx")

    (admin, cro) = _audit_unduhan()
    assert admin[0] == users["ADMIN"]
    assert admin[1]["cakupan"] == "cabang" and admin[1]["format"] == "pdf"
    assert cro[0] == users["CRO"]
    assert cro[1]["cakupan"] == "pribadi" and cro[1]["format"] == "xlsx"
    assert cro[1]["dari"] == hari.isoformat()


def test_cro_tidak_dapat_mencatat_unduhan_laporan_cabang(auth, antrian_siap):
    from app.services.report_service import ReportService

    hari = clock.today()
    auth.logout()
    auth.login("admin", TEST_PASSWORD)
    cabang = ReportService(auth).laporan(hari, hari)

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    with pytest.raises(AuthorizationError):
        ReportService(auth).catat_unduhan(cabang, "pdf")


# ---------------------------------------------------------------------
# Kartu, sorotan, tren, dan laporan pribadi
# ---------------------------------------------------------------------


def _baris(nama, **isi):
    from app.services.report_service import BarisLaporan

    daftar = {k: isi.pop(k) for k in ("_menunggu", "_layanan", "_rating") if k in isi}
    return BarisLaporan(nama, **isi, **daftar)


def _laporan_buatan(**ganti):
    from datetime import date

    from app.services.report_service import Laporan

    nilai = dict(
        dari=date(2026, 10, 1),
        sampai=date(2026, 10, 3),
        total=_baris("Total", terbit=10, dipanggil=10, selesai=8, tidak_hadir=2),
        per_hari=[],
        per_layanan=[],
        per_loket=[],
        per_petugas=[],
    )
    nilai.update(ganti)
    return Laporan(**nilai)


def test_laporan_pribadi_tanpa_kolom_menunggu_dan_prioritas(auth, laporan_hari_ini):
    from app.services.report_service import (
        ReportService,
        angka_prioritas,
        kolom_riwayat,
        kolom_tabel,
    )

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    hari = clock.today()
    pribadi = ReportService(auth).laporan_petugas(hari, hari)

    assert angka_prioritas(pribadi) == []
    for kunci in ("per_hari", "per_layanan", "per_loket"):
        nama = {k[0] for k in kolom_tabel(kunci, pribadi=True)}
        assert not nama & {"Rata-rata menunggu", "Tunggu prioritas", "Prioritas", "Dicabut"}
        assert "Rata-rata pelayanan" in nama
    assert "Lama menunggu" not in {k[0] for k in kolom_riwayat(pribadi=True)}
    # Laporan cabang tidak berubah.
    assert "Rata-rata menunggu" in {k[0] for k in kolom_tabel("per_hari")}
    assert angka_prioritas(laporan_hari_ini)


def test_total_sebelumnya_terisi(laporan_hari_ini):
    sebelum = laporan_hari_ini.total_sebelumnya
    assert sebelum is not None
    assert sebelum.terbit == 0


def test_kartu_ringkasan_nada_dan_perubahan():
    from app.services.report_service import kartu_ringkasan

    laporan = _laporan_buatan(
        total=_baris(
            "Total",
            terbit=12,
            dipanggil=10,
            selesai=8,
            tidak_hadir=2,
            _rating=[5, 5, 4],
            _menunggu=[300.0],
            _layanan=[120.0],
        ),
        total_sebelumnya=_baris(
            "Total",
            terbit=10,
            dipanggil=10,
            selesai=10,
            tidak_hadir=1,
            _rating=[3, 3],
            _menunggu=[600.0],
        ),
    )
    kartu = {k.nama: k for k in kartu_ringkasan(laporan)}

    assert len(kartu) == 6
    assert kartu["Nomor terbit"].perubahan.startswith("▲ 20%")
    assert kartu["Nomor terbit"].nada_perubahan == "netral"
    assert kartu["Selesai dilayani"].perubahan.startswith("▼ 20%")
    assert kartu["Selesai dilayani"].nada_perubahan == "buruk"
    assert kartu["Tidak hadir"].nada == "buruk"  # 20% dari yang dipanggil
    assert kartu["Tidak hadir"].catatan == "20% dari yang dipanggil"
    assert kartu["Rata-rata menunggu"].nada_perubahan == "baik"  # 10 mnt -> 5 mnt
    assert kartu["Kepuasan"].nilai == "4,7 / 5"
    assert kartu["Kepuasan"].nada == "baik"
    assert kartu["Kepuasan"].nada_perubahan == "baik"


def test_kartu_ringkasan_tanpa_pembanding_dan_pribadi():
    from app.services.report_service import kartu_ringkasan

    kartu = kartu_ringkasan(_laporan_buatan(nama_petugas="Cro"))
    nama = [k.nama for k in kartu]
    assert nama[0] == "Nomor ditangani"
    assert "Nomor dipanggil" in nama and "Rata-rata menunggu" not in nama
    assert all(k.perubahan == "" for k in kartu)
    assert kartu[-1].nilai == "—" and kartu[-1].nada == "netral"


def test_sorotan_butuh_pembanding():
    from app.services.report_service import sorotan

    sepi = _laporan_buatan(per_hari=[_baris("2026-10-01", terbit=3)])
    assert sorotan(sepi) == []

    ramai = _laporan_buatan(
        per_hari=[_baris("2026-10-01", terbit=3), _baris("2026-10-02", terbit=9)],
        per_layanan=[_baris("Umum", terbit=9), _baris("Khusus", terbit=3)],
        per_loket=[
            _baris("Loket 1", selesai=3, _layanan=[600.0]),
            _baris("Loket 2", selesai=3, _layanan=[300.0]),
        ],
        per_petugas=[
            _baris("Ani", _rating=[5, 5, 4]),
            _baris("Budi", _rating=[3, 3, 3]),
            _baris("Cici", _rating=[5]),  # survei terlalu sedikit
        ],
    )
    temuan = dict(sorotan(ramai))
    assert "9 nomor" in temuan["Hari tersibuk"]
    assert temuan["Layanan terbanyak"].startswith("Umum")
    assert temuan["Loket tercepat"].startswith("Loket 2")
    assert temuan["Kepuasan tertinggi"].startswith("Ani")

    pribadi = _laporan_buatan(
        nama_petugas="Ani",
        per_loket=ramai.per_loket,
        per_petugas=ramai.per_petugas,
        per_hari=ramai.per_hari,
    )
    assert "Loket tercepat" not in dict(sorotan(pribadi))


def test_deret_tren_harian_dan_mingguan():
    from datetime import date

    from app.services.report_service import deret_tren

    harian = _laporan_buatan(per_hari=[_baris("2026-10-02", selesai=4)])
    judul, deret = deret_tren(harian)
    assert judul == "Nomor selesai per hari"
    assert deret == [("1", 0), ("2", 4), ("3", 0)]

    panjang = _laporan_buatan(
        dari=date(2026, 9, 1),
        sampai=date(2026, 10, 15),
        per_hari=[_baris("2026-09-02", selesai=2), _baris("2026-09-05", selesai=3)],
    )
    judul, deret = deret_tren(panjang)
    assert judul == "Nomor selesai per minggu"
    assert deret[0] == ("1/9", 5)
    assert len(deret) == 7


def test_survei_dihitung_pada_hari_nomor_terbit(laporan_hari_ini):
    (hari,) = laporan_hari_ini.per_hari
    assert hari.terbit == 4 and hari.jumlah_survey == 1


def test_rentang_cepat_tahun_ini_dan_label():
    from datetime import date

    from app.ui.dashboard.report_page import label_rentang, rentang_cepat

    hari = date(2026, 10, 7)
    assert rentang_cepat("tahun_ini", hari) == (date(2026, 1, 1), hari)
    assert label_rentang(date(2026, 10, 1), date(2026, 10, 19)) == "1 Okt – 19 Okt 2026 (19 hari)"
    assert label_rentang(hari, hari) == "7 Okt 2026 (1 hari)"
    assert "2025" in label_rentang(date(2025, 12, 30), date(2026, 1, 2))


def test_halaman_kosong_menyembunyikan_tabel(qtbot, auth, antrian_siap):
    from app.services.report_service import ReportService
    from app.ui.dashboard.report_page import PESAN_KOSONG, LaporanPage

    auth.logout()
    auth.login("supervisor", TEST_PASSWORD)
    halaman = LaporanPage(ReportService(auth), "KANTOR")
    qtbot.addWidget(halaman)
    halaman.show()
    halaman._pakai_rentang("hari_ini")

    assert not halaman.tab.isVisible()
    assert not halaman.grafik.isVisible() and not halaman.sorotan.isVisible()
    assert halaman.kosong.isVisible() and halaman.kosong.text() == PESAN_KOSONG
    assert "dibandingkan dengan" in halaman.label_rentang.text()


def test_halaman_tab_riwayat_saya_hanya_untuk_pribadi(qtbot, auth, laporan_hari_ini):
    from app.services.report_service import ReportService
    from app.ui.dashboard.report_page import LaporanPage

    def judul_kolom(tabel):
        return [
            tabel.horizontalHeaderItem(i).text().replace("\n", " ")
            for i in range(tabel.columnCount())
        ]

    cabang = LaporanPage(ReportService(auth), "KANTOR")
    qtbot.addWidget(cabang)
    assert "riwayat" not in cabang.tabel and "per_petugas" in cabang.tabel
    assert "Rata-rata menunggu" in judul_kolom(cabang.tabel["per_hari"])

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    pribadi = LaporanPage(ReportService(auth), "KANTOR", pribadi=True)
    qtbot.addWidget(pribadi)
    pribadi._pakai_rentang("hari_ini")
    assert list(pribadi.tabel)[:2] == ["per_hari", "riwayat"]
    assert "per_petugas" not in pribadi.tabel
    assert pribadi.tabel["riwayat"].rowCount() == 3
    assert "Rata-rata menunggu" not in judul_kolom(pribadi.tabel["per_hari"])
    assert pribadi.kartu[3].label_nama.text() == "Nomor dipanggil"


def test_item_tabel_diurutkan_menurut_angka_bukan_teks(qtbot):
    from app.ui.dashboard.report_page import ItemUrut

    kecil, besar = ItemUrut("9", 9), ItemUrut("10", 10)
    assert kecil < besar
    assert ItemUrut("", None) < kecil


def test_ekspor_pribadi_tanpa_kolom_menunggu(auth, laporan_hari_ini, tmp_path):
    from openpyxl import load_workbook

    from app.services.report_export import ekspor_excel
    from app.services.report_service import ReportService

    auth.logout()
    auth.login("cro", TEST_PASSWORD)
    hari = clock.today()
    pribadi = ReportService(auth).laporan_petugas(hari, hari)

    buku = load_workbook(ekspor_excel(pribadi, tmp_path / "saya.xlsx"))
    for lembar in buku.worksheets:
        if lembar.title == "Ringkasan":
            continue
        judul = [sel.value for sel in lembar[1]]
        assert not any("menunggu" in str(j).lower() for j in judul if j), lembar.title
        assert "Prioritas dicabut" not in judul
    ringkasan = [baris[0].value for baris in buku["Ringkasan"].iter_rows()]
    assert "Rata-rata tunggu prioritas" not in ringkasan
    assert "Rata-rata menunggu" not in ringkasan


@pytest.mark.parametrize("lebar", [1024, 1280, 1366])
@pytest.mark.parametrize("pribadi", [False, True])
def test_tabel_tanpa_gulir_mendatar_dan_judul_muat(
    qtbot, auth, laporan_hari_ini, lebar, pribadi
):
    from app.services.report_service import ReportService
    from app.ui.dashboard.report_page import LaporanPage

    if pribadi:
        auth.logout()
        auth.login("cro", TEST_PASSWORD)
    halaman = LaporanPage(ReportService(auth), "KANTOR", pribadi=pribadi)
    qtbot.addWidget(halaman)
    halaman.resize(lebar, 700)
    halaman.show()
    halaman._pakai_rentang("hari_ini")

    for kunci, tabel in halaman.tabel.items():
        halaman.tab.setCurrentWidget(tabel)
        qtbot.wait(20)
        assert not tabel.horizontalScrollBar().isVisible(), (kunci, lebar)
        kepala = tabel.horizontalHeader()
        metrik = kepala.fontMetrics()
        jumlah = sum(kepala.sectionSize(c) for c in range(tabel.columnCount()))
        assert jumlah <= tabel.viewport().width()
        for kolom in range(tabel.columnCount()):
            judul = tabel.horizontalHeaderItem(kolom).text()
            terpanjang = max(metrik.horizontalAdvance(b) for b in judul.split("\n"))
            assert terpanjang + 16 <= kepala.sectionSize(kolom), (kunci, judul, lebar)


def test_laporan_menghitung_catatan_identik_per_petugas(
    auth, antrian_siap, cro_queue, terbitkan
):
    from app.services.note_service import FRASA_BAWAAN, NoteService
    from app.services.report_service import ReportService, kolom_tabel

    catatan = NoteService(auth)
    for isi in (
        "Berkas belum lengkap, kembali besok.",
        "BERKAS belum  lengkap, kembali besok.",  # sama setelah dirapikan
        "Menanyakan tagihan bulan ini.",
        FRASA_BAWAAN[0],  # kalimat cepat: tidak dihitung
        FRASA_BAWAAN[0],
    ):
        terbitkan("A")
        antrian = cro_queue.call_next(antrian_siap["loket"][0])
        cro_queue.start_service(antrian.id)
        catatan.simpan(antrian.id, catatan=isi)
        cro_queue.complete(antrian.id)

    auth.logout()
    auth.login("supervisor", TEST_PASSWORD)
    hari = clock.today()
    laporan = ReportService(auth).laporan(hari, hari)

    (petugas,) = laporan.per_petugas
    assert petugas.catatan_bebas == 3
    assert petugas.catatan_identik == 1
    assert petugas.label_catatan_identik == "1 dari 3 (33%)"
    assert "Catatan identik" in [k[0] for k in kolom_tabel("per_petugas")]
    kolom = {k[0]: k for k in kolom_tabel("per_petugas")}["Catatan identik"]
    assert kolom[2](petugas) == 33.3


def test_catatan_identik_kosong_bila_tidak_ada_catatan(laporan_hari_ini):
    (petugas,) = laporan_hari_ini.per_petugas
    assert petugas.catatan_bebas == 0
    assert petugas.label_catatan_identik == "—"
