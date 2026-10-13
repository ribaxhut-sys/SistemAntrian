"""
Ekspor laporan ke PDF (dicetak, ditandatangani) dan Excel (diolah lagi).

Keduanya memakai KOLOM_TABEL yang sama dengan layar, sehingga angka di
kertas tidak pernah berbeda dengan angka yang dilihat supervisor. Excel
menyimpan angka sebagai angka (menit, desimal), bukan teks "9 menit",
supaya dapat langsung dijumlah dan dibuat grafik.

Rincian petugas per jenis layanan masuk ke keduanya; riwayat per nomor
antrian hanya ke Excel, karena setahun penuh bisa ratusan halaman PDF.
"""

from __future__ import annotations

from pathlib import Path

from app.core import clock
from app.services.report_service import (
    JUDUL_TABEL,
    JUDUL_TABEL_PETUGAS_LAYANAN,
    JUDUL_TABEL_RIWAYAT,
    KOLOM_PETUGAS_LAYANAN,
    KOLOM_RIWAYAT,
    KOLOM_TABEL,
    Laporan,
    angka_ringkasan,
)
from app.utils.formatting import tanggal_panjang


def judul_rentang(laporan: Laporan) -> str:
    if laporan.dari == laporan.sampai:
        return tanggal_panjang(laporan.dari)
    return f"{tanggal_panjang(laporan.dari)} – {tanggal_panjang(laporan.sampai)}"


def nama_berkas_bawaan(laporan: Laporan, akhiran: str) -> str:
    return f"laporan-antrian-{laporan.dari:%Y%m%d}-{laporan.sampai:%Y%m%d}.{akhiran}"


def ekspor_pdf(laporan: Laporan, tujuan: Path, kantor: str = "") -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    gaya = getSampleStyleSheet()
    tujuan = Path(tujuan)
    dokumen = SimpleDocTemplate(
        str(tujuan),
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"Laporan Antrian {judul_rentang(laporan)}",
    )
    isi = []
    if kantor:
        isi.append(Paragraph(kantor, gaya["Heading3"]))
    isi.append(Paragraph("Laporan Antrian dan Kepuasan", gaya["Title"]))
    isi.append(Paragraph(judul_rentang(laporan), gaya["Normal"]))
    isi.append(Spacer(1, 6 * mm))

    ringkasan = angka_ringkasan(laporan)
    tabel_ringkas = Table(
        [[nilai for _, nilai in ringkasan], [nama for nama, _ in ringkasan]],
        colWidths=[44 * mm] * len(ringkasan),
    )
    tabel_ringkas.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, 0), 16),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 1), (-1, 1), 8),
                ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#64748B")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    isi.append(tabel_ringkas)

    for kunci, judul in JUDUL_TABEL.items():
        baris = getattr(laporan, kunci)
        if not baris:
            continue
        kolom = KOLOM_TABEL[kunci]
        isi.append(Spacer(1, 6 * mm))
        isi.append(Paragraph(judul, gaya["Heading3"]))
        data = [[k[0] for k in kolom]] + [[k[1](b) for k in kolom] for b in baris]
        tabel = Table(data, repeatRows=1)
        perintah = [
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#CBD5E1")),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        for i, k in enumerate(kolom):
            if k[3]:
                perintah.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))
        tabel.setStyle(TableStyle(perintah))
        isi.append(tabel)

    if laporan.petugas_layanan:
        # Dikelompokkan per petugas: nama hanya di baris pertamanya, lalu
        # baris jumlah yang sama persis dengan tabel Per petugas.
        kolom = KOLOM_PETUGAS_LAYANAN
        isi.append(Spacer(1, 6 * mm))
        isi.append(Paragraph(JUDUL_TABEL_PETUGAS_LAYANAN, gaya["Heading3"]))
        data = [[k[0] for k in kolom]]
        perintah = [
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        for total, rincian in laporan.petugas_layanan:
            for i, baris in enumerate(rincian):
                data.append([k[1](baris) for k in kolom])
                if i:
                    data[-1][0] = ""
            jumlah = [k[1](total) for k in kolom]
            jumlah[0], jumlah[1] = "", "Jumlah"
            data.append(jumlah)
            r = len(data) - 1
            perintah += [
                ("FONTNAME", (1, r), (-1, r), "Helvetica-Bold"),
                ("LINEABOVE", (1, r), (-1, r), 0.25, colors.HexColor("#CBD5E1")),
                ("LINEBELOW", (0, r), (-1, r), 0.5, colors.HexColor("#CBD5E1")),
            ]
        for i, k in enumerate(kolom):
            if k[3]:
                perintah.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))
        tabel = Table(data, repeatRows=1)
        tabel.setStyle(TableStyle(perintah))
        isi.append(tabel)

    isi.append(Spacer(1, 6 * mm))
    isi.append(
        Paragraph(
            f"Dicetak {clock.now():%d-%m-%Y %H:%M} WIB. Menunggu = terbit sampai "
            "dipanggil; pelayanan = mulai dilayani sampai selesai.",
            gaya["Italic"],
        )
    )
    dokumen.build(isi)
    return tujuan


def ekspor_excel(laporan: Laporan, tujuan: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buku = Workbook()
    ringkas = buku.active
    ringkas.title = "Ringkasan"
    ringkas.append(["Laporan Antrian dan Kepuasan"])
    ringkas["A1"].font = Font(bold=True, size=14)
    ringkas.append([judul_rentang(laporan)])
    ringkas.append([])
    for nama, nilai in angka_ringkasan(laporan):
        ringkas.append([nama, nilai])
    ringkas.column_dimensions["A"].width = 32
    ringkas.column_dimensions["B"].width = 16

    tebal = Font(bold=True)
    latar = PatternFill("solid", fgColor="F1F5F9")

    def isi_lembar(judul: str, kolom: list, daftar: list, saring: bool = False) -> None:
        lembar = buku.create_sheet(judul)
        lembar.append([k[0] for k in kolom])
        for sel in lembar[1]:
            sel.font = tebal
            sel.fill = latar
        for baris in daftar:
            lembar.append([k[2](baris) for k in kolom])
        for i, k in enumerate(kolom, start=1):
            huruf = get_column_letter(i)
            lembar.column_dimensions[huruf].width = max(12, len(k[0]) + 4)
            if k[0] == "Tanggal":
                for sel in lembar[huruf][1:]:
                    sel.number_format = "DD-MM-YYYY"
            if k[0].startswith("Jam "):
                for sel in lembar[huruf][1:]:
                    sel.number_format = "HH:MM:SS"
            if k[0].startswith(("Rata-rata", "Lama ")):
                lembar.cell(row=1, column=i).value = f"{k[0]} (menit)"
                lembar.column_dimensions[huruf].width = len(k[0]) + 12
            if k[0] == "Kepuasan":
                for sel in lembar[huruf][1:]:
                    sel.number_format = "0.00"
            if not k[3]:
                for sel in lembar[huruf]:
                    sel.alignment = Alignment(horizontal="left")
        lembar.freeze_panes = "A2"
        if saring and daftar:
            lembar.auto_filter.ref = lembar.dimensions

    for kunci, judul in JUDUL_TABEL.items():
        isi_lembar(judul, KOLOM_TABEL[kunci], getattr(laporan, kunci))
    isi_lembar(
        JUDUL_TABEL_PETUGAS_LAYANAN,
        KOLOM_PETUGAS_LAYANAN,
        [baris for _, rincian in laporan.petugas_layanan for baris in rincian],
        saring=True,
    )
    isi_lembar(JUDUL_TABEL_RIWAYAT, KOLOM_RIWAYAT, laporan.riwayat, saring=True)

    tujuan = Path(tujuan)
    buku.save(tujuan)
    return tujuan
